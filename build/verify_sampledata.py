# -*- coding: utf-8 -*-
"""
modUtils.Test_GenerateSampleData の生成ロジックを Python で忠実に再現し、
実際に生成された設定シートのマスタに対して妥当性を検証する。
（VBA はヘッドレス環境で実行できないため、アルゴリズムの正当性確認に用いる）
"""
import os
import random
import datetime
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WB = os.path.join(ROOT, "人生奪還PJ_資金管理ツール_v1.xlsm")


def load_masters():
    wb = load_workbook(WB)
    cfg = wb["設定"]
    # カテゴリマスタ A3:D46
    cat = []
    r = 3
    while cfg.cell(row=r, column=1).value:
        cat.append((cfg.cell(row=r, column=1).value,
                    cfg.cell(row=r, column=2).value,
                    cfg.cell(row=r, column=3).value))
        r += 1
    majors_order = []
    middles_by_major = {}
    inexp_by_major = {}
    for maj, mid, ie in cat:
        if maj not in majors_order:
            majors_order.append(maj)
            middles_by_major[maj] = []
            inexp_by_major[maj] = ie
        middles_by_major[maj].append(mid)
    payments = []
    r = 3
    while cfg.cell(row=r, column=6).value:
        payments.append(cfg.cell(row=r, column=6).value)
        r += 1
    return majors_order, middles_by_major, inexp_by_major, payments


def generate(majors, middles_by_major, inexp_by_major, payments, n=10):
    today = datetime.date.today()
    first = today.replace(day=1)
    if first.month == 12:
        last = first.replace(year=first.year + 1, month=1, day=1) - datetime.timedelta(days=1)
    else:
        last = first.replace(month=first.month + 1, day=1) - datetime.timedelta(days=1)
    day_count = (last - first).days + 1
    rows = []
    for i in range(n):
        major = majors[i % len(majors)]              # 大分類を巡回選択
        mids = middles_by_major[major]
        middle = random.choice(mids)                 # 中分類はランダム
        pay = payments[i % len(payments)]            # 支払方法を巡回配分
        amount = (random.randint(0, 499) + 1) * 100  # 100〜50,000円
        d = first + datetime.timedelta(days=random.randint(0, day_count - 1))
        rows.append([d, major, middle, f"サンプル{i+1:02d}", amount, pay,
                     inexp_by_major[major], ""])
    return rows


def main():
    majors, middles_by_major, inexp_by_major, payments = load_masters()
    rows = generate(majors, middles_by_major, inexp_by_major, payments)
    print("=== 生成された10件（アルゴリズム再現） ===")
    ok = True
    for row in rows:
        d, maj, mid, desc, amt, pay, ie, _ = row
        v_mid = mid in middles_by_major[maj]
        v_ie = ie == inexp_by_major[maj]
        v_pay = pay in payments
        v_amt = 100 <= amt <= 50000 and amt == int(amt)
        v_date = isinstance(d, datetime.date)
        allok = all([v_mid, v_ie, v_pay, v_amt, v_date])
        ok = ok and allok
        print(f"{d}  {maj}/{mid}  {amt:>6,}円  {pay}  {ie}  "
              f"[中分類整合={v_mid} 収支={v_ie} 支払={v_pay} 金額={v_amt}]")
    print("=== 検証結果:", "全件OK" if ok else "NG", "===")
    # 支払方法の配分
    from collections import Counter
    print("支払方法配分:", dict(Counter(r[5] for r in rows)))
    print("大分類配分:", dict(Counter(r[1] for r in rows)))


if __name__ == "__main__":
    main()
